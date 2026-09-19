import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.model_selection import train_test_split, GridSearchCV
from sklearn.preprocessing import StandardScaler
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import confusion_matrix, roc_curve, auc, accuracy_score, precision_score, recall_score, f1_score
from docx import Document
from docx.shared import Inches
import os

# --- 1. GENERATE MOCK DATA (Classification Task) ---
np.random.seed(42)
n_samples = 300
data = pd.DataFrame({
    'Customer_Age': np.random.randint(18, 70, n_samples),
    'Income': np.random.uniform(30000, 120000, n_samples),
    'Time_on_Site': np.random.uniform(1, 15, n_samples)
})
# Target Variable: 1 if Purchased, 0 if Not. Higher income and time on site increases purchase probability.
probability = (data['Income'] / 120000) * 0.4 + (data['Time_on_Site'] / 15) * 0.6
data['Purchased'] = np.where(probability + np.random.normal(0, 0.1, n_samples) > 0.6, 1, 0)

# --- 2. DATA PREPROCESSING ---
X = data[['Customer_Age', 'Income', 'Time_on_Site']]
y = data['Purchased']

# Train-Test Split (80% Train, 20% Test)
X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)

# Feature Scaling
scaler = StandardScaler()
X_train_scaled = scaler.fit_transform(X_train)
X_test_scaled = scaler.transform(X_test)

# --- 3. MODEL TRAINING & PARAMETER TUNING ---
# We chose Random Forest because it handles non-linear relationships well and is robust to overfitting.
rf_model = RandomForestClassifier(random_state=42)

# Grid Search for Hyperparameter Tuning
param_grid = {
    'n_estimators': [50, 100, 150],
    'max_depth': [None, 5, 10]
}
grid_search = GridSearchCV(estimator=rf_model, param_grid=param_grid, cv=5, scoring='accuracy')
grid_search.fit(X_train_scaled, y_train)

best_model = grid_search.best_estimator_
y_pred = best_model.predict(X_test_scaled)
y_prob = best_model.predict_proba(X_test_scaled)[:, 1]

# --- 4. MODEL EVALUATION ---
acc = accuracy_score(y_test, y_pred)
prec = precision_score(y_test, y_pred)
rec = recall_score(y_test, y_pred)
f1 = f1_score(y_test, y_pred)

# Visual 1: Confusion Matrix
plt.figure(figsize=(5, 4))
cm = confusion_matrix(y_test, y_pred)
sns.heatmap(cm, annot=True, fmt='d', cmap='Blues')
plt.title('Confusion Matrix')
plt.ylabel('Actual')
plt.xlabel('Predicted')
plt.tight_layout()
plt.savefig('confusion_matrix.png')
plt.close()

# Visual 2: ROC Curve
fpr, tpr, thresholds = roc_curve(y_test, y_prob)
roc_auc = auc(fpr, tpr)

plt.figure(figsize=(5, 4))
plt.plot(fpr, tpr, color='darkorange', lw=2, label=f'ROC curve (area = {roc_auc:.2f})')
plt.plot([0, 1], [0, 1], color='navy', lw=2, linestyle='--')
plt.title('Receiver Operating Characteristic (ROC)')
plt.xlabel('False Positive Rate')
plt.ylabel('True Positive Rate')
plt.legend(loc="lower right")
plt.tight_layout()
plt.savefig('roc_curve.png')
plt.close()

# --- 5. BUILD WORD DOCUMENT ---
doc = Document()
doc.add_heading('Machine Learning Algorithms and Model Evaluation', 0)
doc.add_paragraph('Author: Faizhan Ali Syed')

doc.add_heading('1. Algorithm Selection & Justification', level=1)
doc.add_paragraph("Task: Classification (Predicting Customer Purchase).\nAlgorithm: Random Forest Classifier.\nJustification: Random Forest was selected because it is an ensemble learning method that handles non-linear data exceptionally well, reduces the risk of overfitting compared to single decision trees, and does not require extensive data distribution assumptions.")

doc.add_heading('2. Preprocessing & Parameter Tuning', level=1)
doc.add_paragraph("Preprocessing: The data was split into an 80/20 train-test ratio. Features were normalized using StandardScaler to ensure uniform scale for the algorithm.")
doc.add_paragraph(f"Parameter Tuning (Grid Search CV):\nTested Parameters: {param_grid}\nBest Parameters Found: {grid_search.best_params_}")

doc.add_heading('3. Model Evaluation Metrics', level=1)
doc.add_paragraph(f"Accuracy: {acc:.4f} (Overall correctness of the model)")
doc.add_paragraph(f"Precision: {prec:.4f} (Accuracy of positive predictions)")
doc.add_paragraph(f"Recall: {rec:.4f} (Ability to find all positive instances)")
doc.add_paragraph(f"F1-Score: {f1:.4f} (Harmonic mean of precision and recall)")

doc.add_heading('4. Visualizations', level=1)
doc.add_paragraph("Confusion Matrix: Displays the true positives, true negatives, false positives, and false negatives.")
doc.add_picture('confusion_matrix.png', width=Inches(4.5))

doc.add_paragraph("ROC Curve: Illustrates the diagnostic ability of the binary classifier system.")
doc.add_picture('roc_curve.png', width=Inches(4.5))

doc.add_heading('5. Conclusion & Recommendations', level=1)
doc.add_paragraph("The Random Forest model achieved strong predictive performance. For further iterations, I recommend exploring Gradient Boosting algorithms (like XGBoost) to see if sequential tree building improves the F1-score, and testing on a larger, more imbalanced dataset to simulate real-world conditions.")

# --- 6. SAVE DIRECTLY TO DESKTOP FOLDER ---
output_path = r"C:\Users\VICTUS\Desktop\New folder\ML_Model_Evaluation.docx"
doc.save(output_path)
print(f"Done! ML Report saved to {output_path}")